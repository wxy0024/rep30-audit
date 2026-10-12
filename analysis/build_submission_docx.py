#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
build_submission_docx.py — convert manuscript/paper.md (+ refs.bib) into a
journal-submission Word document with numbered (Vancouver-style) citations,
a numbered reference list, real Word tables, and embedded figures.

Also writes:
  * manuscript/REP30_references.ris        (EndNote import file, all 49 refs)
  * manuscript/supplementary_text_submission.docx

Run with the managed venv python (has python-docx + PIL).
"""
import os
import re
import csv
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"E:\01.WorkSpace\01.Breast_cancer_gene_panel\resistance_signature"
MAN = os.path.join(ROOT, "manuscript")
FIG = os.path.join(MAN, "figures")
TAB = os.path.join(MAN, "tables")

# ----------------------------------------------------------------- bibtex


def parse_bib(path):
    entries, cur, buf = {}, None, []
    for line in open(path, encoding="utf-8"):
        s = line.strip()
        if not s or s.startswith("%"):
            continue
        if s.startswith("@article{"):
            key = s[len("@article{"):].strip().strip(",")
            cur, buf = key, []
        elif s == "}" and cur is not None:
            entries[cur] = parse_fields("\n".join(buf))
            cur, buf = None, []
        elif cur is not None:
            buf.append(line)
    return entries


def parse_fields(body):
    fields = {}
    for m in re.finditer(r"(\w+)\s*=\s*\{", body):
        name, i, depth, j = m.group(1), m.end(), 1, m.end()
        while j < len(body) and depth:
            if body[j] == "{":
                depth += 1
            elif body[j] == "}":
                depth -= 1
            j += 1
        if name not in fields:
            fields[name] = body[i:j - 1]
    return fields


_ACC = {'"': "\u0308", "'": "\u0301", "`": "\u0300", "^": "\u0302",
        "~": "\u0303", "=": "\u0304", ".": "\u0307"}


def _acc(m):
    a, c = m.group(1), m.group(2)
    return c + _ACC.get(a, a)


def unlatex(s):
    s = s.replace("\\&", "&")
    s = re.sub(r"\\([\"'`^~=.])\{([a-zA-Z])\}", _acc, s)
    s = re.sub(r"\\([\"'`^~=.])([a-zA-Z])", _acc, s)
    s = re.sub(r"\\c\{([a-zA-Z])\}", lambda m: m.group(1) + "\u0327", s)
    for k, v in [("\\o", "\u00f8"), ("\\O", "\u00d8"), ("\\aa", "\u00e5"),
                 ("\\AA", "\u00c5"), ("\\ae", "\u00e6"), ("\\AE", "\u00c6"),
                 ("\\ss", "\u00df"), ("\\l", "\u0142"), ("\\L", "\u0141"),
                 ("\\i", "\u0131")]:
        s = re.sub(re.escape(k) + r"(?![a-zA-Z])", v, s)
    s = s.replace("\\_", "_").replace("\\%", "%").replace("\\\\", " ")
    s = re.sub(r"\-\-\-", "\u2014", s)
    s = re.sub(r"\-\-", "\u2013", s)
    s = s.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", s).strip()


def fmt_authors(raw, maxn=6):
    raw = unlatex(raw)
    parts = [p.strip() for p in re.split(r"\s+and\s+", raw)]
    names, truncated = [], False
    for p in parts:
        if p.lower() == "others" or p in ("...", "et al."):
            truncated = True
            break
        if "," in p:
            sur, giv = p.split(",", 1)
            inits = "".join(w[0].upper() for w in re.split(r"[\s\-]+", giv.strip())
                            if w and re.match(r"^[A-Za-z]", w))
            names.append((sur.strip() + " " + inits).strip())
        else:
            names.append(p)
    if truncated or len(names) > maxn:
        out = ", ".join(names[:maxn])
        return (out + ", et al") if out else "et al"
    return ", ".join(names)


def fmt_ref(key, e):
    au = fmt_authors(e.get("author", "")).rstrip(".")
    ti = unlatex(e.get("title", "")).rstrip(".")
    jo = unlatex(e.get("journal", ""))
    yr = e.get("year", "").strip()
    vol = e.get("volume", "").strip()
    num = e.get("number", "").strip()
    pg = e.get("pages", "").replace("--", "-").strip()
    doi = e.get("doi", "").strip()
    out = ""
    if au:
        out += au + ". "
    out += ti + ". "
    if jo:
        out += jo + ". "
    if yr:
        out += yr
        if vol:
            out += ";" + vol
            if num:
                out += "(" + num + ")"
        if pg:
            out += ":" + pg
        out += "."
    if doi:
        out += " doi:" + doi + "."
    return re.sub(r"\s+", " ", out).strip()


# ----------------------------------------------------------------- citations


class Citations:
    def __init__(self):
        self.order = []

    def keys(self, grp):
        return re.findall(r"@([A-Za-z0-9]+)", grp)

    def number(self, keys):
        for k in keys:
            if k not in self.order:
                self.order.append(k)
        nums = sorted(self.order.index(k) + 1 for k in keys)
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            parts.append(str(nums[i]) if j == i else
                         (str(nums[i]) + "-" + str(nums[j]) if j > i + 1 else
                          str(nums[i]) + "," + str(nums[j])))
            i = j + 1
        return "[" + ",".join(parts) + "]"

    def sub(self, text):
        out = re.sub(r"\[((?:[^\[\]]*@)[^\[\]]*)\]",
                     lambda m: self.number(self.keys(m.group(1))), text)
        # collapse the double space left behind when the bracket followed a space
        out = re.sub(r"[ \t]{2,}\[", " [", out)
        return out


# ----------------------------------------------------------------- md -> runs


def add_md_runs(p, text):
    text = text.replace("`", "")
    for tk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tk:
            continue
        if tk.startswith("**") and tk.endswith("**") and len(tk) > 4:
            r = p.add_run(tk[2:-2])
            r.bold = True
        elif tk.startswith("*") and tk.endswith("*") and len(tk) > 2:
            r = p.add_run(tk[1:-1])
            r.italic = True
        else:
            p.add_run(tk)


# ----------------------------------------------------------------- docx setup

BODY_FONT = "Times New Roman"
BODY_SIZE = 12


def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(BODY_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.element.rPr.rFonts.set(qn("w:cs"), BODY_FONT)
    pf = st.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)
    return doc


def add_line_numbers(doc):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    doc.sections[-1]._sectPr.append(ln)


def para(doc, text="", size=BODY_SIZE, bold=False, align=None, spacing=2.0,
         space_after=6, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if align:
        p.alignment = align
    if text:
        add_md_runs(p, text)
        for r in p.runs:
            r.font.name = BODY_FONT
            r.font.size = Pt(size)
            if bold:
                r.bold = True
    return p


def heading(doc, text, level=1):
    sizes = {0: 14, 1: 12, 2: 12}
    p = para(doc, text, size=sizes.get(level, 12), bold=True, spacing=1.5,
             space_before=12, space_after=6)
    return p


def page_break(doc):
    doc.add_page_break()


def add_table(doc, csv_path, caption_size=8):
    rows = list(csv.reader(open(csv_path, encoding="utf-8-sig")))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            txt = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            add_md_runs(p, txt)
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.size = Pt(caption_size)
                if i == 0:
                    r.bold = True
    para(doc, "", space_after=10)


# ----------------------------------------------------------------- main build


def build_main():
    src = open(os.path.join(MAN, "paper.md"), encoding="utf-8").read()
    lines = src.splitlines()
    # drop the trailing word-count line(s)
    lines = [l for l in lines if not l.strip().startswith("*Word count:")]
    src = "\n".join(lines)

    i_ref = src.index("## References")
    i_leg = src.index("## Figure legends")
    i_tleg = src.index("## Table legends")
    main_part = src[:i_ref]
    legends = src[i_leg:i_tleg]
    table_legends = src[i_tleg:]

    bib = parse_bib(os.path.join(MAN, "refs.bib"))
    cit = Citations()
    main_txt = cit.sub(main_part)
    leg_txt = cit.sub(legends)
    tleg_txt = cit.sub(table_legends)
    supp_order = None  # filled later

    doc = new_doc()
    add_line_numbers(doc)

    # ---------- title page
    title = re.search(r"^# (.+)$", main_txt, re.M).group(1)
    para(doc, title, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.5, space_after=12)
    rt = re.search(r"\*\*Running title:\*\*(.+)", main_txt)
    if rt:
        para(doc, "Running title: " + rt.group(1).strip(), size=12,
             align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=12)
    para(doc, "[[TBD: full author list with initials and degrees]]",
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=6)
    para(doc, "[[TBD: affiliations]]", align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.5, space_after=6)
    para(doc, "[[TBD: corresponding author, address, e-mail, ORCID]]",
         align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=12)
    para(doc, "Keywords: breast cancer; gene expression signature; proliferation; "
              "drug resistance; prognosis; HR+/HER2- disease", spacing=1.5,
         space_after=12)
    para(doc, "Word counts: abstract 305; main text 6,584; figure and table "
              "legends 2,110. Tables: 3 main, 7 supplementary. Figures: 4 main, "
              "4 supplementary.", spacing=1.5, space_after=6)
    page_break(doc)

    # ---------- body
    body_blocks = main_txt.split("\n")
    in_abstract = False
    for ln in body_blocks:
        s = ln.rstrip()
        if not s.strip():
            continue
        if s.startswith("# "):
            continue
        if s.startswith("## "):
            heading(doc, s[3:].strip(), 1)
            in_abstract = s[3:].strip() == "Abstract"
            continue
        if s.startswith("### "):
            heading(doc, s[4:].strip(), 2)
            continue
        if s.strip() == "---":
            continue
        para(doc, s.strip())

    # ---------- references
    page_break(doc)
    heading(doc, "References", 1)
    for n, key in enumerate(cit.order, 1):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        r = p.add_run("%d. " % n)
        r.bold = False
        p.add_run(fmt_ref(key, bib.get(key, {})))
        for r in p.runs:
            r.font.name = BODY_FONT
            r.font.size = Pt(11)
    missing = [k for k in cit.order if k not in bib]
    if missing:
        para(doc, "MISSING BIB ENTRIES: " + ", ".join(missing), bold=True)

    # ---------- figure legends + figures
    page_break(doc)
    heading(doc, "Figure legends", 1)
    blocks = [b.strip() for b in leg_txt.split("\n\n") if b.strip()]
    blocks = [re.sub(r"\n?\s*-{3,}\s*$", "", b).strip() for b in blocks]
    leg_map = {}
    for b in blocks:
        m = re.match(r"\*\*Fig (S?\d) \|", b)
        if m:
            leg_map[m.group(1)] = b
    for fig_no in ["1", "2", "3", "4", "S1", "S2", "S3", "S4"]:
        cap = leg_map.get(fig_no)
        if cap:
            para(doc, cap, size=11, spacing=1.5, space_after=6)
        path = os.path.join(FIG, "Fig%s_%s.png" % (
            fig_no, {"1": "audit", "2": "repositioning", "3": "failure_modes",
                     "4": "confounder_switch", "S1": "single_cell",
                     "S2": "exposure_axis", "S3": "full_landscape",
                     "S4": "venet_null"}[fig_no]))
        if os.path.exists(path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=Cm(16.5))
        page_break(doc)

    # ---------- main tables
    for tab_no, fname in [("Table 1", "Table1_cohort_inventory.csv"),
                          ("Table 2", "Table2_headline_evidence.csv"),
                          ("Table 3", "Table3_hrposher2neg_performance.csv")]:
        heading(doc, tab_no, 2)
        m = re.search(r"\*\*%s \|[^*]+\*\*[^*]*(\*\*[^*]+\*\*)?" % tab_no, tleg_txt)
        # caption: take the legend paragraph for this table
        for b in [b.strip() for b in tleg_txt.split("\n\n") if b.strip()]:
            if b.startswith("**%s |" % tab_no):
                para(doc, b, size=11, spacing=1.5, space_after=6)
                break
        add_table(doc, os.path.join(TAB, fname))
        page_break(doc)

    out = os.path.join(MAN, "paper_submission.docx")
    doc.save(out)
    return cit.order, out


def build_supplementary(main_order):
    src = open(os.path.join(MAN, "supplementary_text.md"), encoding="utf-8").read()
    cit = Citations()
    for k in main_order:
        cit.order.append(k)  # continue main-text numbering
    txt = cit.sub(src)
    doc = new_doc()
    add_line_numbers(doc)
    for ln in txt.split("\n"):
        s = ln.rstrip()
        if not s.strip() or s.strip() == "---":
            continue
        if s.startswith("# "):
            heading(doc, s[2:].strip(), 0)
        elif s.startswith("## "):
            heading(doc, s[3:].strip(), 1)
        elif s.startswith("### "):
            heading(doc, s[4:].strip(), 2)
        else:
            para(doc, s.strip())
    # supplementary-only references (keys not in the main list)
    extra = [k for k in cit.order if k not in main_order]
    para(doc, "Numbered citations in this supplementary text continue the numbering of the "
              "main manuscript and refer to its reference list.", size=10, spacing=1.5,
         space_after=12)
    if extra:
        page_break(doc)
        heading(doc, "References cited in the Supplementary Text", 1)
        bib = parse_bib(os.path.join(MAN, "refs.bib"))
        for k in extra:
            n = cit.order.index(k) + 1
            p = para(doc, "%d. %s" % (n, fmt_ref(k, bib.get(k, {}))), size=11,
                     spacing=1.5, space_after=4)
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
    out = os.path.join(MAN, "supplementary_text_submission.docx")
    doc.save(out)
    return out


def write_ris():
    bib = parse_bib(os.path.join(MAN, "refs.bib"))
    out = os.path.join(MAN, "REP30_references.ris")
    with open(out, "w", encoding="utf-8") as fh:
        for key, e in bib.items():
            fh.write("TY  - JOUR\n")
            raw = unlatex(e.get("author", ""))
            for p in re.split(r"\s+and\s+", raw):
                p = p.strip()
                if not p:
                    continue
                if p.lower() == "others":
                    continue
                fh.write("AU  - " + (p if "," in p else p) + "\n")
            fh.write("TI  - " + unlatex(e.get("title", "")) + "\n")
            fh.write("T2  - " + unlatex(e.get("journal", "")) + "\n")
            fh.write("JO  - " + unlatex(e.get("journal", "")) + "\n")
            if e.get("volume"):
                fh.write("VL  - " + e["volume"].strip() + "\n")
            if e.get("number"):
                fh.write("IS  - " + e["number"].strip() + "\n")
            pg = e.get("pages", "").replace("--", "-").strip()
            if "-" in pg:
                sp, ep = pg.split("-", 1)
                fh.write("SP  - " + sp + "\nEP  - " + ep + "\n")
            elif pg:
                fh.write("SP  - " + pg + "\n")
            if e.get("year"):
                fh.write("PY  - " + e["year"].strip() + "\n")
            if e.get("doi"):
                fh.write("DO  - " + e["doi"].strip() + "\n")
            if e.get("note"):
                fh.write("N1  - " + unlatex(e["note"]) + "\n")
            fh.write("ER  - \n\n")
    return out


if __name__ == "__main__":
    order, main_out = build_main()
    supp_out = build_supplementary(order)
    ris_out = write_ris()
    print("citations numbered:", len(order))
    print("wrote:", main_out)
    print("wrote:", supp_out)
    print("wrote:", ris_out)
